import shutil, subprocess, json, sys
from openpyxl import load_workbook
import datetime

BASE = "Campaign_Control_System.xlsx"
RECALC = "/root/.claude/skills/synced/46159b45-ccc8-440d-ab2d-516922183b8f_b76b9256-fe07-4209-932c-3f32f4509599/xlsx/scripts/recalc.py"

TODAY = datetime.date(2026, 8, 8)

PCT_OBS = [0.492, 0.577, 0.614, 0.641, 0.663, 0.679, 0.693, 0.710]  # ages 0-7; age>=8 -> 1.0


def true_roas_overrides(true_roas, n_days=14, spend=5000):
    """Build {age_offset: (spend, revenue)} so that age-adjusted ROAS recovers true_roas
    regardless of the age mix in the window (mirrors how the sample data itself was built)."""
    out = {}
    for age in range(n_days):
        pct = PCT_OBS[age] if age < 8 else 1.0
        revenue = round(spend * true_roas * pct, 2)
        out[age] = (spend, revenue)
    return out


scenarios = {
    "01_baseline": {},
    "02_safety_pause": {
        # Nuke the last 3 days of Daily Data to well below breakeven
        "daily_overrides": {  # row offset from bottom (0=most recent) -> (spend, revenue)
            0: (5000, 4000), 1: (5000, 4200), 2: (5000, 4300),
        }
    },
    "03_severe_pause_primary_only": {
        # Uniform true ROAS 195% (below breakeven 200.1%) across all ages 0-13.
        # Fast window (ages 0-2) also averages ~195%, i.e. above the 180.1% safety-pause
        # trigger, so this exercises the *confirmed, primary-window* pause path, not the
        # fast-window safety net.
        "daily_overrides": true_roas_overrides(1.95),
    },
    "04_tighten": {
        "daily_overrides": true_roas_overrides(2.05),  # aim ~205%: between breakeven and normal band
    },
    "05_loosen": {
        "daily_overrides": true_roas_overrides(2.45),  # above 235% normal-band upper edge
        "inputs": {"R_CUR_BUDGET": 15000},  # generous budget -> not capped
    },
    "06_budget_increase": {
        "daily_overrides": true_roas_overrides(2.45),
        "inputs": {"R_CUR_BUDGET": 5000},  # spend==budget => capped
    },
    "07_paused_hold": {
        "inputs": {"R_STATUS": "Paused", "R_LAST_CHANGE_DATE": datetime.date(2026, 8, 7),
                   "R_LAST_CHANGE_TYPE": "Pause"},
    },
    "08_paused_reenable": {
        "inputs": {"R_STATUS": "Paused", "R_LAST_CHANGE_DATE": datetime.date(2026, 8, 3),
                   "R_LAST_CHANGE_TYPE": "Pause"},
    },
    "09_wait_hold": {
        "inputs": {"R_LAST_CHANGE_DATE": datetime.date(2026, 8, 6),
                   "R_LAST_CHANGE_TYPE": "Target ROAS Change"},
    },
    "10_data_insufficient": {
        "orders_fixed": 1,  # 21 sample days x 1 order = 21 total, below the 30-order confidence floor
                            # in BOTH the 14-day primary and (all-available) 30-day fallback windows
    },
    "11_season_start": {
        "inputs": {"R_SEASON_ACTIVE": "Y", "R_SEASON_DIR": "Up",
                   "R_SEASON_START": datetime.date(2026, 8, 8), "R_SEASON_DUR": 10,
                   "R_SEASON_SIZE": 0.15, "R_SEASON_BASELINE": 2.20},
    },
    "12_season_mid": {
        "inputs": {"R_SEASON_ACTIVE": "Y", "R_SEASON_DIR": "Up",
                   "R_SEASON_START": datetime.date(2026, 8, 3), "R_SEASON_DUR": 10,
                   "R_SEASON_SIZE": 0.15, "R_SEASON_BASELINE": 2.20},
    },
    "13_season_end": {
        "inputs": {"R_SEASON_ACTIVE": "Y", "R_SEASON_DIR": "Up",
                   "R_SEASON_START": datetime.date(2026, 7, 30), "R_SEASON_DUR": 10,
                   "R_SEASON_SIZE": 0.15, "R_SEASON_BASELINE": 2.20},
    },
}

INPUT_ROWS = {
    "R_TODAY": 4, "R_STATUS": 5, "R_CUR_TARGET": 6, "R_CUR_BUDGET": 7,
    "R_LAST_CHANGE_DATE": 8, "R_LAST_CHANGE_TYPE": 9,
    "R_SEASON_ACTIVE": 12, "R_SEASON_DIR": 13, "R_SEASON_START": 14,
    "R_SEASON_DUR": 15, "R_SEASON_SIZE": 16, "R_SEASON_BASELINE": 17,
}

results = {}
for name, cfg in scenarios.items():
    fn = f"scenario_{name}.xlsx"
    shutil.copy(BASE, fn)
    wb = load_workbook(fn, data_only=False)
    ws_in = wb["Inputs"]
    ws_dd = wb["Daily Data"]

    for key, val in cfg.get("inputs", {}).items():
        ws_in.cell(row=INPUT_ROWS[key], column=2, value=val)

    if "daily_overrides" in cfg:
        # Daily Data rows: DD_FIRST=3, ascending oldest->newest, so "offset from bottom" 0 = row for date TODAY
        n_sample = 21
        DD_FIRST = 3
        last_row = DD_FIRST + n_sample - 1  # row for most recent date (today)
        for offset, (spend, revenue) in cfg["daily_overrides"].items():
            row = last_row - offset
            ws_dd.cell(row=row, column=2, value=spend)
            ws_dd.cell(row=row, column=3, value=revenue)

    if "orders_fixed" in cfg:
        n_sample = 21
        DD_FIRST = 3
        for i in range(n_sample):
            row = DD_FIRST + i
            ws_dd.cell(row=row, column=4, value=cfg["orders_fixed"])

    wb.save(fn)
    r = subprocess.run(["python3", RECALC, fn, "90"], capture_output=True, text=True)
    out = json.loads(r.stdout)
    if out.get("status") != "success":
        results[name] = {"error": out}
        continue

    wbv = load_workbook(fn, data_only=True)
    d = wbv["Decision"]
    results[name] = {
        "case": wbv["Engine"]["B48"].value,
        "action": d["B18"].value,
        "value": d["B19"].value,
        "cadence": d["B21"].value,
        "wait": d["B22"].value,
        "adjroas_used": d["B7"].value,
        "fast_adjroas": d["B8"].value,
        "window_used": d["B6"].value,
        "orders_used": d["B9"].value,
        "budget_capped": d["B15"].value,
    }

for name, r in results.items():
    print("=" * 70)
    print(name)
    for k, v in r.items():
        print(f"  {k}: {v}")
